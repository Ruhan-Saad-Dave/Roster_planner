"""Export service generating formatted multi-sheet Excel (.xlsx) and printable PDF schedules."""
import io
import json
from datetime import timedelta, date
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

# Excel Color Palette
COLOR_HEADER_BG = "1E293B"       # Slate 800
COLOR_HEADER_TEXT = "FFFFFF"
COLOR_MORN_BG = "DBEAFE"         # Blue 100
COLOR_MORN_TEXT = "1E40AF"       # Blue 800
COLOR_EVE_BG = "FEF3C7"          # Amber 100
COLOR_EVE_TEXT = "92400E"        # Amber 800
COLOR_NIGHT_BG = "E0E7FF"        # Indigo 100
COLOR_NIGHT_TEXT = "3730A3"      # Indigo 800
COLOR_OFF_BG = "F8FAFC"          # Slate 50
COLOR_OFF_TEXT = "94A3B8"        # Slate 400
COLOR_DO_BG = "E2E8F0"           # Slate 200 (Dedicated Off)
COLOR_DO_TEXT = "334155"         # Slate 700
COLOR_COFF_BG = "FEF3C7"         # Amber 100 (Comp Off)
COLOR_COFF_TEXT = "B45309"       # Amber 700
COLOR_VAC_BG = "EDE9FE"          # Violet 100 (Leave/Vacation)
COLOR_VAC_TEXT = "6D28D9"        # Violet 700
COLOR_DEFICIT_BG = "FEE2E2"      # Red 100
COLOR_DEFICIT_TEXT = "991B1B"    # Red 800
COLOR_FLOAT_BORDER = "F59E0B"    # Amber 500

def get_shift_fill_and_font(shift_code: Optional[str], is_floating: bool = False):
    border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )
    if is_floating:
        border = Border(
            left=Side(style="medium", color=COLOR_FLOAT_BORDER),
            right=Side(style="medium", color=COLOR_FLOAT_BORDER),
            top=Side(style="medium", color=COLOR_FLOAT_BORDER),
            bottom=Side(style="medium", color=COLOR_FLOAT_BORDER)
        )

    code = (shift_code or "").upper()
    if not code or code == "OFF":
        return PatternFill("solid", fgColor=COLOR_OFF_BG), Font(name="Calibri", size=9, color=COLOR_OFF_TEXT), border
    elif code in ["DO", "DAY_OFF", "DAY OFF"]:
        return PatternFill("solid", fgColor=COLOR_DO_BG), Font(name="Calibri", size=9, bold=True, color=COLOR_DO_TEXT), border
    elif "COFF" in code or "COMP" in code:
        return PatternFill("solid", fgColor=COLOR_COFF_BG), Font(name="Calibri", size=9, bold=True, color=COLOR_COFF_TEXT), border
    elif any(k in code for k in ["VAC", "SICK", "LEAVE", "TRAIN"]):
        return PatternFill("solid", fgColor=COLOR_VAC_BG), Font(name="Calibri", size=9, bold=True, color=COLOR_VAC_TEXT), border
    elif "MORN" in code:
        return PatternFill("solid", fgColor=COLOR_MORN_BG), Font(name="Calibri", size=9, bold=True, color=COLOR_MORN_TEXT), border
    elif "EVE" in code:
        return PatternFill("solid", fgColor=COLOR_EVE_BG), Font(name="Calibri", size=9, bold=True, color=COLOR_EVE_TEXT), border
    elif "NIGHT" in code:
        return PatternFill("solid", fgColor=COLOR_NIGHT_BG), Font(name="Calibri", size=9, bold=True, color=COLOR_NIGHT_TEXT), border
    return PatternFill("solid", fgColor="E2E8F0"), Font(name="Calibri", size=9, color="334155"), border

def generate_excel_export(
    plan: Any,
    assignments: List[Any],
    deficits: List[Any],
    employees: List[Any],
    departments: List[Any],
    shifts: List[Any],
    leaves: Optional[List[Any]] = None,
    custom_settings: Optional[Dict[str, Any]] = None,
    view: str = "both"
) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    horizon_days = (plan.end_date - plan.start_date).days + 1
    dates = [plan.start_date + timedelta(days=t) for t in range(horizon_days)]

    emp_map = {e.id: e for e in employees}
    dept_map = {d.id: d for d in departments}
    shift_map = {s.id: s for s in shifts}

    # Custom labels dictionary if configured by manager
    custom_labels = {}
    if custom_settings and "custom_status_codes" in custom_settings:
        try:
            val = custom_settings["custom_status_codes"]
            custom_labels = json.loads(val) if isinstance(val, str) else val
        except Exception:
            pass

    label_off = custom_labels.get("OFF", "OFF")
    label_coff = custom_labels.get("COFF", "COFF")
    label_vac = custom_labels.get("VAC", "VAC")
    label_do = custom_labels.get("DO", "DO")

    # Map assignments: (emp_id, date) -> assign
    assign_by_emp_date = {(a.employee_id, a.date): a for a in assignments}

    # Map approved leaves: (emp_id, date) -> leave
    leave_by_emp_date = {}
    if leaves:
        for l in leaves:
            cur = l.start_date
            while cur <= l.end_date:
                leave_by_emp_date[(l.employee_id, cur)] = l
                cur += timedelta(days=1)

    # Precalculate employee weekly workload to identify overtime earned comp offs
    emp_weekly_hours = {}
    emp_weekly_shifts = {}
    for a in assignments:
        if a.shift_type_id in shift_map:
            wk_key = (a.employee_id, a.date.isocalendar()[:2])
            dur = shift_map[a.shift_type_id].duration_hours or 8.0
            emp_weekly_hours[wk_key] = emp_weekly_hours.get(wk_key, 0.0) + dur
            emp_weekly_shifts[wk_key] = emp_weekly_shifts.get(wk_key, 0) + 1

    # ==========================================
    # TAB 1: Staff-Centric Schedule
    # ==========================================
    if view in ["staff_centric", "both"]:
        ws_staff = wb.create_sheet(title="Staff Schedule")
        ws_staff.views.sheetView[0].showGridLines = True

        # Header Title
        ws_staff.cell(row=1, column=1, value=f"LocalRoster v2 - {plan.name} (Staff-Centric Monthly Schedule)")
        ws_staff.cell(row=1, column=1).font = Font(name="Calibri", size=14, bold=True, color="1E293B")
        ws_staff.cell(row=2, column=1, value=f"Horizon: {plan.start_date} to {plan.end_date} | Solver Status: {plan.solver_status} | Staff Count: {len(employees)} | Total Shifts: {len(assignments)}")
        ws_staff.cell(row=2, column=1).font = Font(name="Calibri", size=9, italic=True, color="64748B")

        # Table Headers
        headers = ["Code", "Employee Name", "Home Dept", "Dedicated Shift", "Backup Floats"]
        for d in dates:
            headers.append(d.strftime("%b %d\n(%a)"))
        headers.extend(["Worked Shifts", "Hours Sched", "Target (Mo)", "Variance", "Regular DO", "Comp Offs", "Leaves", "Floats"])

        for col_idx, h in enumerate(headers, start=1):
            cell = ws_staff.cell(row=4, column=col_idx, value=h)
            cell.font = Font(name="Calibri", size=9, bold=True, color=COLOR_HEADER_TEXT)
            cell.fill = PatternFill("solid", fgColor=COLOR_HEADER_BG)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws_staff.row_dimensions[4].height = 28

        # Data Rows
        row_idx = 5
        for emp in employees:
            ws_staff.cell(row=row_idx, column=1, value=emp.employee_code).alignment = Alignment(horizontal="center")
            ws_staff.cell(row=row_idx, column=2, value=emp.name).font = Font(name="Calibri", size=9, bold=True)
            
            home_dept = dept_map.get(emp.primary_department_id)
            ws_staff.cell(row=row_idx, column=3, value=home_dept.code if home_dept else "").alignment = Alignment(horizontal="center")

            # Dedicated shift
            ded_shift = shift_map.get(emp.preferred_shift_type_id) if emp.preferred_shift_type_id else None
            ded_shift_str = ded_shift.name if ded_shift else "Dynamic / Flexible"
            ws_staff.cell(row=row_idx, column=4, value=ded_shift_str).alignment = Alignment(horizontal="center")

            # Secondary floating eligibilities
            sec_eligs = [el for el in (emp.eligibilities or []) if el.preference_rank > 1 and el.can_float]
            sec_codes = [dept_map[el.department_id].code for el in sec_eligs if el.department_id in dept_map]
            ws_staff.cell(row=row_idx, column=5, value=", ".join(sec_codes) if sec_codes else "-").alignment = Alignment(horizontal="center")

            shifts_worked = 0
            hours_worked = 0.0
            float_count = 0
            do_count = 0
            coff_count = 0
            leave_count = 0

            emp_fixed_days = []
            if emp.fixed_off_days:
                try:
                    emp_fixed_days = [int(x.strip()) for x in emp.fixed_off_days.split(",") if x.strip().isdigit()]
                except Exception:
                    pass

            for d_idx, cur_date in enumerate(dates, start=6):
                cell = ws_staff.cell(row=row_idx, column=d_idx)
                if (emp.id, cur_date) in assign_by_emp_date:
                    a = assign_by_emp_date[(emp.id, cur_date)]
                    s = shift_map.get(a.shift_type_id)
                    d_code = dept_map.get(a.department_id).code if a.department_id in dept_map else ""
                    val = s.code if s else "WORK"
                    if a.is_floating:
                        val += f" ({d_code}*)"
                        float_count += 1
                    if a.is_pinned:
                        val += " [P]"

                    cell.value = val
                    shifts_worked += 1
                    if s:
                        hours_worked += s.duration_hours

                    fill, font, border = get_shift_fill_and_font(s.code if s else "MORN", a.is_floating)
                    cell.fill = fill
                    cell.font = font
                    cell.border = border
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                elif (emp.id, cur_date) in leave_by_emp_date:
                    l = leave_by_emp_date[(emp.id, cur_date)]
                    ltype = (l.leave_type or "").upper()
                    lnotes = (l.notes or "").upper()
                    if ltype in ["COFF", "COMP_OFF"] or "COFF" in lnotes:
                        cell.value = f"{label_coff} [P]"
                        coff_count += 1
                        fill, font, border = get_shift_fill_and_font("COFF")
                    elif ltype in ["DAY_OFF", "DO"] or "DAY OFF" in lnotes:
                        cell.value = f"{label_do} [P]"
                        do_count += 1
                        fill, font, border = get_shift_fill_and_font("DO")
                    else:
                        code = label_vac if ltype == "VACATION" else ltype[:3]
                        cell.value = f"{code} [P]"
                        leave_count += 1
                        fill, font, border = get_shift_fill_and_font("VAC")
                    cell.fill = fill
                    cell.font = font
                    cell.border = border
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                elif cur_date.weekday() in emp_fixed_days:
                    cell.value = f"{label_do} (Reg)"
                    do_count += 1
                    fill, font, border = get_shift_fill_and_font("DO")
                    cell.fill = fill
                    cell.font = font
                    cell.border = border
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                else:
                    wk_key = (emp.id, cur_date.isocalendar()[:2])
                    w_hours = emp_weekly_hours.get(wk_key, 0.0)
                    w_shifts = emp_weekly_shifts.get(wk_key, 0)
                    contract_hours = emp.max_hours_week or 40.0
                    is_solved = plan.status in ["SOLVED", "PUBLISHED"]

                    if is_solved and (w_hours >= contract_hours or w_shifts >= 5):
                        cell.value = f"{label_coff} (Earned)"
                        coff_count += 1
                        fill, font, border = get_shift_fill_and_font("COFF")
                    else:
                        cell.value = label_off
                        fill, font, border = get_shift_fill_and_font("OFF")

                    cell.fill = fill
                    cell.font = font
                    cell.border = border
                    cell.alignment = Alignment(horizontal="center", vertical="center")

            # Totals
            target_hours = round(emp.max_hours_week * (horizon_days / 7.0), 1)
            variance = round(hours_worked - target_hours, 1)

            base_stat_col = len(dates) + 6
            ws_staff.cell(row=row_idx, column=base_stat_col, value=shifts_worked).alignment = Alignment(horizontal="center")
            ws_staff.cell(row=row_idx, column=base_stat_col + 1, value=f"{hours_worked:.1f}h").alignment = Alignment(horizontal="center")
            ws_staff.cell(row=row_idx, column=base_stat_col + 2, value=f"{target_hours:.1f}h").alignment = Alignment(horizontal="center")

            var_cell = ws_staff.cell(row=row_idx, column=base_stat_col + 3, value=f"{'+' if variance > 0 else ''}{variance:.1f}h")
            var_cell.alignment = Alignment(horizontal="center")
            if variance > 8.0:
                var_cell.font = Font(name="Calibri", size=9, bold=True, color="DC2626")
            elif variance < -8.0:
                var_cell.font = Font(name="Calibri", size=9, color="D97706")

            ws_staff.cell(row=row_idx, column=base_stat_col + 4, value=do_count).alignment = Alignment(horizontal="center")
            ws_staff.cell(row=row_idx, column=base_stat_col + 5, value=coff_count).alignment = Alignment(horizontal="center")
            ws_staff.cell(row=row_idx, column=base_stat_col + 6, value=leave_count).alignment = Alignment(horizontal="center")
            ws_staff.cell(row=row_idx, column=base_stat_col + 7, value=float_count).alignment = Alignment(horizontal="center")

            row_idx += 1

        # Add Legend Block at Bottom
        leg_row = row_idx + 2
        ws_staff.cell(row=leg_row, column=1, value="Color & Code Legend:").font = Font(name="Calibri", size=10, bold=True)
        legends = [
            ("MORN", "Morning Shift (07:00-15:30)", COLOR_MORN_BG, COLOR_MORN_TEXT),
            ("EVE", "Evening Shift (15:00-23:30)", COLOR_EVE_BG, COLOR_EVE_TEXT),
            ("NIGHT", "Night Shift (23:00-07:30)", COLOR_NIGHT_BG, COLOR_NIGHT_TEXT),
            ("DO", "Dedicated Regular Day Off / Locked DO", COLOR_DO_BG, COLOR_DO_TEXT),
            ("COFF", "Compensatory Off (Earned Overtime / Locked)", COLOR_COFF_BG, COLOR_COFF_TEXT),
            ("VAC", "Approved Vacation / Leave", COLOR_VAC_BG, COLOR_VAC_TEXT),
            ("OFF", "Scheduled Rest Day", COLOR_OFF_BG, COLOR_OFF_TEXT),
        ]
        for l_idx, (code, desc, bg, fg) in enumerate(legends):
            c_cell = ws_staff.cell(row=leg_row + 1 + l_idx, column=2, value=code)
            c_cell.fill = PatternFill("solid", fgColor=bg)
            c_cell.font = Font(name="Calibri", size=9, bold=True, color=fg)
            c_cell.alignment = Alignment(horizontal="center")
            ws_staff.cell(row=leg_row + 1 + l_idx, column=3, value=desc).font = Font(name="Calibri", size=9)

        # Set Column Widths
        ws_staff.column_dimensions["A"].width = 11
        ws_staff.column_dimensions["B"].width = 20
        ws_staff.column_dimensions["C"].width = 12
        ws_staff.column_dimensions["D"].width = 18
        ws_staff.column_dimensions["E"].width = 15
        for col in range(6, len(dates) + 6):
            col_letter = openpyxl.utils.get_column_letter(col)
            ws_staff.column_dimensions[col_letter].width = 11
        for col in range(len(dates) + 6, len(headers) + 1):
            col_letter = openpyxl.utils.get_column_letter(col)
            ws_staff.column_dimensions[col_letter].width = 13

    # ==========================================
    # TAB 2: Slot-Centric Schedule
    # ==========================================
    if view in ["slot_centric", "both"]:
        ws_slot = wb.create_sheet(title="Coverage by Slot")
        ws_slot.views.sheetView[0].showGridLines = True

        ws_slot.cell(row=1, column=1, value=f"LocalRoster v2 - {plan.name} (Department & Slot Coverage View)")
        ws_slot.cell(row=1, column=1).font = Font(name="Calibri", size=14, bold=True, color="1E293B")

        slot_headers = ["Date", "Day", "Department", "Shift", "Assigned Staff", "Total Staff", "Deficit Status"]
        for col_idx, h in enumerate(slot_headers, start=1):
            cell = ws_slot.cell(row=3, column=col_idx, value=h)
            cell.font = Font(name="Calibri", size=9, bold=True, color=COLOR_HEADER_TEXT)
            cell.fill = PatternFill("solid", fgColor=COLOR_HEADER_BG)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        slot_row = 4
        assign_by_slot = {}
        for a in assignments:
            key = (a.date, a.department_id, a.shift_type_id)
            assign_by_slot.setdefault(key, []).append(a)

        deficit_by_slot = {(d.date, d.department_id, d.shift_type_id): d for d in (deficits or [])}

        for cur_date in dates:
            for dept in departments:
                for shift in shifts:
                    key = (cur_date, dept.id, shift.id)
                    slot_assigns = assign_by_slot.get(key, [])
                    deficit_entry = deficit_by_slot.get(key)
                    if slot_assigns or deficit_entry:
                        worker_names = []
                        for sa in slot_assigns:
                            e_name = emp_map[sa.employee_id].name if sa.employee_id in emp_map else f"Emp {sa.employee_id}"
                            if sa.is_floating:
                                e_name += " (Floater*)"
                            worker_names.append(e_name)

                        ws_slot.cell(row=slot_row, column=1, value=cur_date.strftime("%Y-%m-%d"))
                        ws_slot.cell(row=slot_row, column=2, value=cur_date.strftime("%a"))
                        ws_slot.cell(row=slot_row, column=3, value=dept.name)
                        ws_slot.cell(row=slot_row, column=4, value=shift.name)
                        ws_slot.cell(row=slot_row, column=5, value=", ".join(worker_names) if worker_names else "Uncovered")
                        ws_slot.cell(row=slot_row, column=6, value=len(worker_names)).alignment = Alignment(horizontal="center")

                        def_cell = ws_slot.cell(row=slot_row, column=7)
                        if deficit_entry:
                            def_cell.value = f"SHORTAGE (-{deficit_entry.shortage_count})"
                            def_cell.font = Font(name="Calibri", size=9, bold=True, color=COLOR_DEFICIT_TEXT)
                            def_cell.fill = PatternFill("solid", fgColor=COLOR_DEFICIT_BG)
                        else:
                            def_cell.value = "Adequate"
                            def_cell.font = Font(name="Calibri", size=9, color="16A34A")
                        def_cell.alignment = Alignment(horizontal="center")

                        slot_row += 1

        for col in ws_slot.columns:
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws_slot.column_dimensions[col_letter].width = 18

    # ==========================================
    # TAB 3: Workload & Ergonomics Audit
    # ==========================================
    ws_audit = wb.create_sheet(title="Workload Audit")
    ws_audit.views.sheetView[0].showGridLines = True
    ws_audit.cell(row=1, column=1, value=f"LocalRoster v2 - Cross-Check Audit Scorecard ({plan.name})")
    ws_audit.cell(row=1, column=1).font = Font(name="Calibri", size=14, bold=True, color="1E293B")

    audit_headers = ["Code", "Name", "Department", "Dedicated Shift", "Contract (Wk)", "Target (Mo)", "Scheduled", "Variance", "Shifts", "DOs", "COFFs", "Floats"]
    for col_idx, h in enumerate(audit_headers, start=1):
        cell = ws_audit.cell(row=3, column=col_idx, value=h)
        cell.font = Font(name="Calibri", size=9, bold=True, color=COLOR_HEADER_TEXT)
        cell.fill = PatternFill("solid", fgColor=COLOR_HEADER_BG)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    audit_row = 4
    for emp in employees:
        emp_assigns = [a for a in assignments if a.employee_id == emp.id]
        total_shifts = len(emp_assigns)
        total_hours = sum(shift_map[a.shift_type_id].duration_hours for a in emp_assigns if a.shift_type_id in shift_map)
        floats = sum(1 for a in emp_assigns if a.is_floating)

        target_month_hours = round(emp.max_hours_week * (horizon_days / 7.0), 1)
        variance = round(total_hours - target_month_hours, 1)

        ded_shift = shift_map.get(emp.preferred_shift_type_id) if emp.preferred_shift_type_id else None

        ws_audit.cell(row=audit_row, column=1, value=emp.employee_code).alignment = Alignment(horizontal="center")
        ws_audit.cell(row=audit_row, column=2, value=emp.name)
        ws_audit.cell(row=audit_row, column=3, value=dept_map.get(emp.primary_department_id).name if emp.primary_department_id in dept_map else "")
        ws_audit.cell(row=audit_row, column=4, value=ded_shift.name if ded_shift else "Flexible")
        ws_audit.cell(row=audit_row, column=5, value=f"{emp.max_hours_week}h").alignment = Alignment(horizontal="center")
        ws_audit.cell(row=audit_row, column=6, value=f"{target_month_hours}h").alignment = Alignment(horizontal="center")
        ws_audit.cell(row=audit_row, column=7, value=f"{total_hours:.1f}h").alignment = Alignment(horizontal="center")

        var_cell = ws_audit.cell(row=audit_row, column=8, value=f"{'+' if variance > 0 else ''}{variance:.1f}h")
        var_cell.alignment = Alignment(horizontal="center")
        if variance > 8.0:
            var_cell.font = Font(name="Calibri", size=9, bold=True, color="DC2626")
        elif variance < -8.0:
            var_cell.font = Font(name="Calibri", size=9, color="D97706")

        ws_audit.cell(row=audit_row, column=9, value=total_shifts).alignment = Alignment(horizontal="center")
        
        # Count DOs and COFFs
        emp_fixed_days = [int(x.strip()) for x in emp.fixed_off_days.split(",") if x.strip().isdigit()] if emp.fixed_off_days else []
        dos = sum(1 for d in dates if (emp.id, d) not in assign_by_emp_date and d.weekday() in emp_fixed_days)
        ws_audit.cell(row=audit_row, column=10, value=dos).alignment = Alignment(horizontal="center")

        coffs = sum(1 for d in dates if (emp.id, d) in leave_by_emp_date and "COFF" in (leave_by_emp_date[(emp.id, d)].leave_type or "").upper())
        ws_audit.cell(row=audit_row, column=11, value=coffs).alignment = Alignment(horizontal="center")
        ws_audit.cell(row=audit_row, column=12, value=floats).alignment = Alignment(horizontal="center")

        audit_row += 1

    for col in ws_audit.columns:
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws_audit.column_dimensions[col_letter].width = 16

    # ==========================================
    # TAB 4: Deficits & Shortages
    # ==========================================
    ws_def = wb.create_sheet(title="Deficits & Bottlenecks")
    ws_def.views.sheetView[0].showGridLines = True
    ws_def.cell(row=1, column=1, value=f"LocalRoster v2 - Deficit & Shortage Diagnostics ({plan.name})")
    ws_def.cell(row=1, column=1).font = Font(name="Calibri", size=14, bold=True, color="1E293B")

    def_headers = ["Date", "Day", "Department", "Shift", "Shortage Count", "Root Cause Reason Code"]
    for col_idx, h in enumerate(def_headers, start=1):
        cell = ws_def.cell(row=3, column=col_idx, value=h)
        cell.font = Font(name="Calibri", size=9, bold=True, color=COLOR_HEADER_TEXT)
        cell.fill = PatternFill("solid", fgColor=COLOR_HEADER_BG)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    def_row = 4
    if deficits:
        for d in deficits:
            ws_def.cell(row=def_row, column=1, value=d.date.strftime("%Y-%m-%d"))
            ws_def.cell(row=def_row, column=2, value=d.date.strftime("%a"))
            ws_def.cell(row=def_row, column=3, value=dept_map.get(d.department_id).name if d.department_id in dept_map else "")
            ws_def.cell(row=def_row, column=4, value=shift_map.get(d.shift_type_id).name if d.shift_type_id in shift_map else "")
            short_cell = ws_def.cell(row=def_row, column=5, value=d.shortage_count)
            short_cell.font = Font(name="Calibri", size=9, bold=True, color=COLOR_DEFICIT_TEXT)
            short_cell.fill = PatternFill("solid", fgColor=COLOR_DEFICIT_BG)
            short_cell.alignment = Alignment(horizontal="center")
            ws_def.cell(row=def_row, column=6, value=d.reason_code)
            def_row += 1
    else:
        ws_def.cell(row=4, column=1, value="Clean Roster: 0 staffing deficits detected. Full demand satisfied.")
        ws_def.cell(row=4, column=1).font = Font(name="Calibri", size=10, bold=True, color="16A34A")

    for col in ws_def.columns:
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws_def.column_dimensions[col_letter].width = 20

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def generate_pdf_export(
    plan: Any,
    assignments: List[Any],
    employees: List[Any],
    departments: List[Any],
    shifts: List[Any],
    leaves: Optional[List[Any]] = None,
    custom_settings: Optional[Dict[str, Any]] = None
) -> bytes:
    """Generate high-density, vector-rendered printable PDF schedule covering the complete month."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        rightMargin=16, leftMargin=16,
        topMargin=16, bottomMargin=16
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleStyle",
        parent=styles["Heading1"],
        fontSize=14,
        leading=16,
        textColor=colors.HexColor("#0F172A")
    )
    meta_style = ParagraphStyle(
        "MetaStyle",
        parent=styles["Normal"],
        fontSize=8,
        textColor=colors.HexColor("#64748B")
    )
    subhead_style = ParagraphStyle(
        "SubHeadStyle",
        parent=styles["Normal"],
        fontSize=9,
        leading=11,
        fontName="Helvetica-Bold",
        textColor=colors.HexColor("#1E293B")
    )

    elements = []
    elements.append(Paragraph(f"<b>LocalRoster v2</b> — {plan.name} (Monthly Schedule)", title_style))
    elements.append(Paragraph(
        f"Horizon: {plan.start_date} to {plan.end_date} | Solver: {plan.solver_status} | Staff: {len(employees)} | Shifts: {len(assignments)} | Deficits: {plan.total_deficits}",
        meta_style
    ))
    elements.append(Spacer(1, 8))

    horizon_days = (plan.end_date - plan.start_date).days + 1
    all_dates = [plan.start_date + timedelta(days=t) for t in range(horizon_days)]
    assign_by_emp_date = {(a.employee_id, a.date): a for a in assignments}
    shift_map = {s.id: s for s in shifts}
    dept_map = {d.id: d for d in departments}

    leave_by_emp_date = {}
    if leaves:
        for l in leaves:
            cur = l.start_date
            while cur <= l.end_date:
                leave_by_emp_date[(l.employee_id, cur)] = l
                cur += timedelta(days=1)

    # Split dates into chunks of max 16 days to ensure clean printing on landscape page
    chunk_size = 16 if horizon_days > 16 else horizon_days
    date_chunks = [all_dates[i:i + chunk_size] for i in range(0, horizon_days, chunk_size)]

    for chunk_idx, dates in enumerate(date_chunks):
        if chunk_idx > 0:
            elements.append(PageBreak())
            elements.append(Paragraph(f"<b>LocalRoster v2</b> — {plan.name} (Part {chunk_idx + 1}: {dates[0]} to {dates[-1]})", title_style))
            elements.append(Spacer(1, 8))
        else:
            if len(date_chunks) > 1:
                elements.append(Paragraph(f"Part 1: {dates[0]} to {dates[-1]}", subhead_style))
                elements.append(Spacer(1, 4))

        table_data = [["Staff", "Dept", "Ded"] + [d.strftime("%d\n%a") for d in dates] + ["Wk"]]
        for emp in employees:
            dept_code = dept_map.get(emp.primary_department_id).code if emp.primary_department_id in dept_map else ""
            ded_shift_code = shift_map.get(emp.preferred_shift_type_id).code[:2] if emp.preferred_shift_type_id in shift_map else "-"
            
            row = [emp.name[:15], dept_code[:4], ded_shift_code]
            worked = 0

            emp_fixed_days = []
            if emp.fixed_off_days:
                try:
                    emp_fixed_days = [int(x.strip()) for x in emp.fixed_off_days.split(",") if x.strip().isdigit()]
                except Exception:
                    pass

            for d in dates:
                if (emp.id, d) in assign_by_emp_date:
                    a = assign_by_emp_date[(emp.id, d)]
                    s = shift_map.get(a.shift_type_id)
                    code = s.code[:2] if s else "W"
                    if a.is_floating:
                        code += "*"
                    row.append(code)
                    worked += 1
                elif (emp.id, d) in leave_by_emp_date:
                    l = leave_by_emp_date[(emp.id, d)]
                    ltype = (l.leave_type or "").upper()
                    if "COFF" in ltype:
                        row.append("CO")
                    elif "DAY_OFF" in ltype or "DO" in ltype:
                        row.append("DO")
                    else:
                        row.append("VA")
                elif d.weekday() in emp_fixed_days:
                    row.append("DO")
                else:
                    row.append("-")

            row.append(str(worked))
            table_data.append(row)

        pdf_table = Table(table_data)
        pdf_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("ALIGN", (0, 1), (0, -1), "LEFT"),
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ]))
        elements.append(pdf_table)
        elements.append(Spacer(1, 8))

    # Add Shift Legend & Summary on last page
    elements.append(Spacer(1, 6))
    elements.append(Paragraph(
        "<b>Legend:</b> MO = Morning | EV = Evening | NI = Night | * = Floating Cross-Dept | DO = Dedicated Day Off | CO = Compensatory Off | VA = Vacation | - = Rest Day",
        meta_style
    ))

    doc.build(elements)
    return buffer.getvalue()
